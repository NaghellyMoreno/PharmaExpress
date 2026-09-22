"""Carga y normaliza fuentes del proyecto para generación.

Este módulo convierte archivos de texto, documentos Office, PDFs, Excalidraw e
imágenes en un formato uniforme con texto e imágenes base64 para que el motor
pueda enviarlos al modelo con contexto multimodal.
"""

from __future__ import annotations

import base64
import json
import mimetypes
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader


@dataclass
class SourceDocument:
    """Representa una fuente con texto e imágenes asociadas."""

    name: str
    text: str = ""
    images: list[tuple[str, str]] = field(default_factory=list)
    # images: [(mime_type, base64_data)]


def _read_docx(path: Path) -> str:
    """Extrae texto de un documento DOCX manteniendo párrafos y tablas."""
    doc = Document(path)
    chunks: list[str] = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            chunks.append(text)

    for table_index, table in enumerate(doc.tables, start=1):
        chunks.append(f"\n[TABLA {table_index}]")
        for row in table.rows:
            cells = [cell.text.replace("\n", " / ").strip() for cell in row.cells]
            chunks.append(" | ".join(cells))

    return "\n".join(chunks)


def _read_pdf(path: Path) -> str:
    """Extrae texto de un PDF page por page para conservar contexto del documento."""
    reader = PdfReader(str(path))
    chunks: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        chunks.append(f"\n[PÁGINA {i}]\n{text}")
    return "\n".join(chunks)


def _read_excalidraw(path: Path) -> str:
    """Lee textos de una exportación Excalidraw para usarlos como contexto textual."""
    data = json.loads(path.read_text(encoding="utf-8"))
    texts: list[str] = []
    for element in data.get("elements", []):
        if element.get("type") == "text":
            value = (element.get("text") or "").strip()
            if value:
                texts.append(value)
    return "\n".join(texts)


def _read_image(path: Path) -> SourceDocument:
    """Carga una imagen como fuente visual y la transforma a base64."""
    mime, _ = mimetypes.guess_type(path.name)
    mime = mime or "image/png"
    data = base64.b64encode(path.read_bytes()).decode("ascii")
    return SourceDocument(name=path.name, images=[(mime, data)])


def load_source(path_like: str | Path) -> SourceDocument:
    """Carga un archivo fuente y devuelve su texto e imágenes en formato uniforme."""
    path = Path(path_like)
    if not path.exists():
        raise FileNotFoundError(path)

    suffix = path.suffix.lower()

    if suffix in {".txt", ".md", ".csv"}:
        return SourceDocument(name=path.name, text=path.read_text(encoding="utf-8"))
    if suffix == ".docx":
        return SourceDocument(name=path.name, text=_read_docx(path))
    if suffix == ".pdf":
        return SourceDocument(name=path.name, text=_read_pdf(path))
    if suffix == ".excalidraw":
        return SourceDocument(name=path.name, text=_read_excalidraw(path))
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        return _read_image(path)

    raise ValueError(f"Formato no soportado: {suffix}")


def load_sources(paths: list[str]) -> list[SourceDocument]:
    """Carga una lista de fuentes y devuelve una colección normalizada."""
    return [load_source(p) for p in paths]
